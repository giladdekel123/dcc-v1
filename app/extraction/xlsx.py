from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook

from app.extraction.segment import Segment

ROWS_PER_SEGMENT = 50


def _text(value) -> str:
    """Cell value as text; whole-day datetimes (Excel dates) as ISO dates."""
    if isinstance(value, datetime) and value.time() == time(0):
        return value.date().isoformat()
    return str(value)


def extract_xlsx(path: Path) -> tuple[list[Segment], int]:
    """One segment per sheet, split every 50 rows."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        segments = []
        for sheet in workbook.worksheets:
            rows = [(n, " | ".join(_text(v) for v in row if v is not None))
                    for n, row in enumerate(sheet.iter_rows(values_only=True), 1)]
            rows = [(n, text) for n, text in rows if text]
            for i in range(0, len(rows), ROWS_PER_SEGMENT):
                chunk = rows[i:i + ROWS_PER_SEGMENT]
                locator = f'sheet "{sheet.title}", rows {chunk[0][0]}-{chunk[-1][0]}'
                segments.append(Segment(locator, "\n".join(text for _, text in chunk)))
        return segments, len(workbook.worksheets)
    finally:
        workbook.close()
