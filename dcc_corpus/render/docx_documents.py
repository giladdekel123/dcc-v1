"""Word documents: RFI forms, letters, meeting minutes and submittals."""

from pathlib import Path

from docx import Document
from docx.shared import Pt

from dcc_corpus.render.common import RenderContext, fixed_datetime, fmt_date, normalise_ooxml


def _new(ctx: RenderContext) -> Document:
    d = Document()
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(10)
    props = d.core_properties
    props.title = f"{ctx.doc['doc_code']} {ctx.rev} {ctx.doc['title']}"
    props.author = ctx.doc["owner"]
    props.last_modified_by = ctx.doc["owner"]
    props.subject = ctx.project["project"]["name"]
    props.created = props.modified = fixed_datetime(ctx.issued)
    props.revision = 1
    return d


def _save(d: Document, path: Path, ctx: RenderContext) -> None:
    d.save(str(path))
    normalise_ooxml(path, ctx.issued)


def _header(d: Document, ctx: RenderContext, heading: str) -> None:
    d.add_paragraph(ctx.org_name(ctx.code.originator)).runs[0].bold = True
    d.add_heading(heading, level=1)
    _table(d, [
        ["Project", ctx.project["project"]["name"]],
        ["Document", ctx.doc["doc_code"]],
        ["Revision", ctx.rev],
        ["Status", ctx.status_text()],
        ["Date", fmt_date(ctx.issued)],
    ])


def _table(d: Document, rows: list[list[str]], header: bool = False) -> None:
    table = d.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            if header and r == 0:
                cell.paragraphs[0].runs[0].bold = True
    d.add_paragraph()


def render_rfi(path: Path, ctx: RenderContext) -> None:
    c = ctx.content
    d = _new(ctx)
    _header(d, ctx, "Request for Information")
    _table(d, [
        ["RFI subject", c["subject"]],
        ["Raised by", ctx.person_line(c["raised_by"])],
        ["Addressed to", ctx.person_line(c["addressed_to"])],
        ["Response required by", fmt_date(c["response_required_by"])],
        ["Programme impact", c["programme_impact"]],
        ["References", ", ".join(c["references"])],
    ])
    d.add_heading("Query", level=2)
    for p in c["question"]:
        d.add_paragraph(p.strip())
    d.add_heading("Response", level=2)
    d.add_paragraph("(To be completed by the designer.)")
    _save(d, path, ctx)


def render_letter(path: Path, ctx: RenderContext) -> None:
    c = ctx.content
    d = _new(ctx)
    d.add_paragraph(ctx.org_name(ctx.code.originator)).runs[0].bold = True
    to = ctx.person(c["to"])
    d.add_paragraph(f"{c['to']}\n{to['title']}\n{ctx.org_name(to['org'])}")
    d.add_paragraph(f"Our ref: {ctx.doc['doc_code']}\nDate: {fmt_date(ctx.issued)}")
    d.add_paragraph(f"Dear {c['to'].split()[0]},")
    d.add_paragraph(c["subject"]).runs[0].bold = True
    for p in c["paragraphs"]:
        d.add_paragraph(p.strip())
    d.add_paragraph("Yours sincerely,")
    d.add_paragraph(ctx.person_line(c["signed_by"]))
    if c.get("cc"):
        d.add_paragraph("cc: " + "; ".join(ctx.person_line(n) for n in c["cc"]))
    _save(d, path, ctx)


def render_minutes(path: Path, ctx: RenderContext) -> None:
    c = ctx.content
    d = _new(ctx)
    _header(d, ctx, f"Progress Meeting No. {c['meeting_no']} - Minutes")
    _table(d, [
        ["Meeting date", fmt_date(ctx.issued)],
        ["Location", c["location"]],
        ["Chair", ctx.person_line(c["chair"])],
        ["Present", "; ".join(ctx.person_line(n) for n in c["attendees"])],
        ["Apologies", "; ".join(ctx.person_line(n) for n in c.get("apologies", []))],
    ])
    rows = [["Item", "Subject", "Discussion", "Action", "Owner", "Due"]]
    for item in c["items"]:
        rows.append([
            item["ref"], item["title"], item["discussion"].strip(),
            item.get("action", "-"),
            _initials(item["owner"]) if "owner" in item else "-",
            item["due"].strftime("%d/%m/%Y") if "due" in item else "-",
        ])
    _table(d, rows, header=True)
    d.add_paragraph("Initials: " + "; ".join(f"{_initials(n)} = {n}" for n in c["attendees"] + c.get("apologies", [])))
    _save(d, path, ctx)


def render_submittal(path: Path, ctx: RenderContext) -> None:
    c = ctx.content
    d = _new(ctx)
    _header(d, ctx, "Submittal")
    _table(d, [
        ["Title", ctx.doc["title"]],
        ["Submitted by", ctx.person_line(c["submitted_by"])],
        ["For review by", ctx.person_line(c["for_review_by"])],
        ["Specification", c["spec_clause"]],
    ])
    for section in c["sections"]:
        d.add_heading(section["heading"], level=2)
        for p in section["paragraphs"]:
            d.add_paragraph(p.strip())
    _save(d, path, ctx)


def _initials(name: str) -> str:
    return "".join(part[0] for part in name.split())
