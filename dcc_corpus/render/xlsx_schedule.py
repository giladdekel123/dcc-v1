"""Pile schedule spreadsheet."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from dcc_corpus.render.common import RenderContext, fixed_datetime, fmt_date, normalise_ooxml, pile_layout

COLUMNS = ["Pile ref", "Easting", "Northing", "Existing ground (mAOD)", "Cut-off level (mAOD)",
           "Toe level (mAOD)", "Length (m)", "Diameter (mm)"]


def render(path: Path, ctx: RenderContext) -> None:
    g = ctx.content["geometry"]
    wb = Workbook()
    wb.properties.creator = ctx.doc["owner"]
    wb.properties.lastModifiedBy = ctx.doc["owner"]
    wb.properties.title = f"{ctx.doc['doc_code']} {ctx.rev} {ctx.doc['title']}"
    wb.properties.created = wb.properties.modified = fixed_datetime(ctx.issued)

    ws = wb.active
    ws.title = "Pile Schedule"
    header = [
        ("Project", ctx.project["project"]["name"]),
        ("Title", ctx.doc["title"]),
        ("Document", ctx.doc["doc_code"]),
        ("Revision", ctx.rev),
        ("Status", ctx.status_text()),
        ("Date", fmt_date(ctx.issued)),
    ]
    for label, value in header:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(COLUMNS)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    length = round(g["cutoff_level"] - g["toe_level"], 2)
    for p in pile_layout(ctx.seed, g["piles"]):
        ws.append([p["ref"], p["easting"], p["northing"], p["ground_level"],
                   g["cutoff_level"], g["toe_level"], length, g["pile_diameter_mm"]])
    for col, width in zip("ABCDEFGH", [10, 14, 14, 22, 20, 17, 11, 13]):
        ws.column_dimensions[col].width = width

    history = wb.create_sheet("Revision history")
    history.append(["Rev", "Date", "Description"])
    for rev, when, description in ctx.history:
        history.append([rev, when.strftime("%d/%m/%Y"), description])

    wb.save(str(path))
    normalise_ooxml(path, ctx.issued)
