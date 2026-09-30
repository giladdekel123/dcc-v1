"""Shared spreadsheet layout: document header block and revision history sheet."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from dcc_corpus.render.common import RenderContext, fixed_datetime, fmt_date, normalise_ooxml


def new_workbook(ctx: RenderContext, sheet_title: str) -> Workbook:
    wb = Workbook()
    wb.properties.creator = ctx.doc["owner"]
    wb.properties.lastModifiedBy = ctx.doc["owner"]
    wb.properties.title = f"{ctx.doc['doc_code']} {ctx.rev} {ctx.doc['title']}"
    wb.properties.created = wb.properties.modified = fixed_datetime(ctx.issued)
    wb.active.title = sheet_title
    return wb


def write_header(ws, ctx: RenderContext, rev_shown: str | None = None) -> None:
    """Project/document header rows. rev_shown lets a file state a revision other than the register."""
    header = [
        ("Project", ctx.project["project"]["name"]),
        ("Title", ctx.doc["title"]),
        ("Document", ctx.doc["doc_code"]),
        ("Revision", rev_shown or ctx.rev),
        ("Status", ctx.status_text()),
        ("Date", fmt_date(ctx.issued)),
    ]
    for label, value in header:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])


def bold_row(ws, values: list) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)


def save(wb: Workbook, path: Path, ctx: RenderContext) -> None:
    history = wb.create_sheet("Revision history")
    history.append(["Rev", "Date", "Description"])
    for rev, when, description in ctx.history:
        history.append([rev, when.strftime("%d/%m/%Y"), description])
    wb.save(str(path))
    normalise_ooxml(path, ctx.issued)
