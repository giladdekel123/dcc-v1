"""The generated corpus is reproducible, matches the committed copy, and its register is consistent."""

import csv
import hashlib
from collections import Counter
from pathlib import Path

import docx
import openpyxl
import pytest

from dcc_corpus import naming
from dcc_corpus.generate import generate

COMMITTED = Path(__file__).resolve().parent.parent / "corpus"


def tree_hashes(root: Path) -> dict[str, str]:
    return {
        p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(root.rglob("*")) if p.is_file()
    }


def read_csv(root: Path, name: str) -> list[dict]:
    with open(root / name, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


@pytest.fixture(scope="module")
def fresh(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("corpus")
    generate(out)
    return out


def test_generation_is_reproducible(fresh, tmp_path):
    generate(tmp_path)
    assert tree_hashes(tmp_path) == tree_hashes(fresh)


def test_committed_corpus_matches_fresh_generation(fresh):
    assert COMMITTED.is_dir(), "run: python -m dcc_corpus.generate"
    assert tree_hashes(COMMITTED) == tree_hashes(fresh)


def test_files_exist_with_matching_hashes(fresh):
    files = read_csv(fresh, "register/revision_files.csv")
    assert len(files) == 18
    for row in files:
        data = (fresh / row["storage_path"]).read_bytes()
        assert row["storage_path"] == f"{row['original_location']}/{row['filename']}"
        assert int(row["size_bytes"]) == len(data)
        assert row["sha256"] == hashlib.sha256(data).hexdigest()
    on_disk = {p.relative_to(fresh).as_posix() for p in (fresh / naming.PROJECT_ROOT).rglob("*") if p.is_file()}
    assert on_disk == {row["storage_path"] for row in files}


def test_one_primary_file_per_revision(fresh):
    revisions = {(r["doc_code"], r["rev_code"]) for r in read_csv(fresh, "register/revisions.csv")}
    primaries = Counter((f["doc_code"], f["rev_code"]) for f in read_csv(fresh, "register/revision_files.csv")
                        if f["copy_role"] == "primary")
    assert set(primaries) == revisions and set(primaries.values()) == {1}


def test_register_values_follow_the_convention(fresh):
    documents = {d["doc_code"]: d for d in read_csv(fresh, "register/documents.csv")}
    for code, d in documents.items():
        parsed = naming.parse_doc_code(code)
        assert (parsed.wbs, parsed.discipline, parsed.doc_type, parsed.originator) == (
            d["wbs_code"], d["discipline_code"], d["doc_type_code"], d["originator_code"])

    revision_dates = {(r["doc_code"], r["rev_code"]): r["revision_date"] for r in read_csv(fresh, "register/revisions.csv")}
    for s in read_csv(fresh, "register/revision_status.csv"):
        assert naming.status_allowed(documents[s["doc_code"]]["doc_type_code"], s["status_code"])
        assert s["effective_date"] >= revision_dates[(s["doc_code"], s["rev_code"])]

    legacy = {(c["doc_code"], c["rev_code"]) for c in read_csv(fresh, "ground_truth/planted_cases.csv")
              if c["case"] == "legacy_filename"}
    for f in read_csv(fresh, "register/revision_files.csv"):
        parsed = naming.parse_filename(f["filename"])
        if (f["doc_code"], f["rev_code"]) in legacy:
            assert parsed is None
        else:
            assert (str(parsed.doc_code), parsed.rev) == (f["doc_code"], f["rev_code"])


def test_links_and_planted_cases_resolve(fresh):
    documents = {d["doc_code"] for d in read_csv(fresh, "register/documents.csv")}
    revisions = {(r["doc_code"], r["rev_code"]) for r in read_csv(fresh, "register/revisions.csv")}

    def resolves(code, rev):
        return code in documents and (not rev or (code, rev) in revisions)

    links = read_csv(fresh, "links.csv")
    assert len(links) == 10
    assert all(resolves(l["from_doc_code"], l["from_rev_code"]) and resolves(l["to_doc_code"], l["to_rev_code"])
               for l in links)
    cases = read_csv(fresh, "ground_truth/planted_cases.csv")
    assert all(resolves(c["doc_code"], c["rev_code"]) for c in cases)
    assert {c["case"] for c in cases} >= {"correct_not_latest", "titleblock_discrepancy", "legacy_filename", "file_copy"}


def test_superseded_files_are_filed_under_99_superseded(fresh):
    superseded = {(f["doc_code"], f["rev_code"]) for f in read_csv(fresh, "register/revision_files.csv")
                  if "/99 Superseded/" in f["storage_path"]}
    assert superseded == {
        ("KVL-ADM-410-DR-S-0102", "P02"), ("KVL-ADM-410-DR-S-0102", "C01"),
        ("KVL-ADM-410-DR-S-0110", "C01"), ("KVL-BCS-410-CT-S-0001", "P01"),
    }


def test_rendered_content_is_readable(fresh):
    files = {(f["doc_code"], f["rev_code"], f["copy_role"]): fresh / f["storage_path"]
             for f in read_csv(fresh, "register/revision_files.csv")}
    rfi = docx.Document(files[("KVL-CGC-410-RI-S-0012", "P01", "primary")])
    text = "\n".join(p.text for p in rfi.paragraphs) + "\n".join(
        c.text for t in rfi.tables for row in t.rows for c in row.cells)
    assert "soft grey clay" in text and "OP - Open" in text

    ws = openpyxl.load_workbook(files[("KVL-ADM-410-SC-S-0003", "C02", "primary")])["Pile Schedule"]
    toe_levels = {row[5] for row in ws.iter_rows(min_row=9, values_only=True)}
    assert toe_levels == {-11.15}

    primary = files[("KVL-ADM-410-DR-S-0102", "C02", "primary")].read_bytes()
    assert files[("KVL-ADM-410-DR-S-0102", "C02", "copy")].read_bytes() == primary
